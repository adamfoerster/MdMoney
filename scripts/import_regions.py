#!/usr/bin/env python3
"""One-time importer: converts the Regions spreadsheet (PalmBayHouse.xlsx) into markdown notes.

Each yearly sheet becomes one `.md` file per expense row, in the same frontmatter schema the app
uses for the Nubank account (English month keys jan..dec + `-paid`, `type`, `conta`, `year`,
`projected`). Run once to seed the vault; afterwards the app owns the files.

Usage:
    python3 scripts/import_regions.py <xlsx_path> <vault_root> [--account Regions] [--category casa]

It writes into <vault_root>/<account>/ and never overwrites an existing file (skips it instead),
so re-running is safe.
"""
import argparse
import os
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required: pip3 install --user openpyxl")

MONTH_KEYS = ["jan", "feb", "mar", "apr", "may", "jun",
              "jul", "aug", "sep", "oct", "nov", "dec"]

# Spreadsheet columns: A=name, B=Year(total, ignored), C=Projetado/Ano, D..O = Jan..Dec.
COL_NAME = 0
COL_PROJECTED = 2
COL_FIRST_MONTH = 3


def fmt_amount(value):
    """Match the app's formatter: integer when whole, else up to two decimals, trimming zeros."""
    if value is None:
        return ""
    cents = round(abs(float(value)) * 100)
    sign = "-" if value < 0 and cents != 0 else ""
    whole, rem = divmod(cents, 100)
    if rem == 0:
        return f"{sign}{whole}"
    if rem % 10 == 0:
        return f"{sign}{whole}.{rem // 10}"
    return f"{sign}{whole}.{rem:02d}"


def infer_type(month_values):
    non_zero = [v for v in month_values if v not in (None, 0)]
    distinct = {round(float(v), 2) for v in non_zero}
    if len(non_zero) == 0:
        return "eventual"
    if len(non_zero) == 1:
        return "eventual"
    if len(distinct) == 1:
        return "recurring-fixed"
    return "recurring-variable"


def sanitize_filename(name):
    return re.sub(r'[\\/:*?"<>|]', "-", name).strip()


def build_note(title, account, category, year, projected, month_values):
    kind = infer_type(month_values)
    lines = ["---"]
    lines.append(f"title: {title}")
    lines.append(f"category: {category}")
    lines.append(f"conta: {account}")
    lines.append(f"year: {year}")
    lines.append(f"type: {kind}")
    lines.append("subStatus: Active")
    if projected not in (None, "", 0):
        lines.append(f"projected: {fmt_amount(projected)}")
    for key, value in zip(MONTH_KEYS, month_values):
        lines.append(f"{key}: {fmt_amount(value)}")
        lines.append(f"{key}-paid: false")
    lines.append("---")
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("vault_root")
    ap.add_argument("--account", default="Regions")
    ap.add_argument("--category", default="casa")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    out_dir = os.path.join(args.vault_root, args.account)
    os.makedirs(out_dir, exist_ok=True)

    written, skipped = 0, 0
    for ws in wb.worksheets:
        year = None
        if re.fullmatch(r"\d{4}", ws.title.strip()):
            year = int(ws.title.strip())
        if year is None:
            continue  # skip non-year sheets (e.g. scratch "Sheet1")

        for row in ws.iter_rows(values_only=True):
            if not row or row[COL_NAME] in (None, "", "Total", "Year"):
                continue
            name = str(row[COL_NAME]).strip()
            projected = row[COL_PROJECTED] if len(row) > COL_PROJECTED else None
            month_values = []
            for i in range(12):
                idx = COL_FIRST_MONTH + i
                month_values.append(row[idx] if len(row) > idx else None)

            file_name = f"{sanitize_filename(name)} - {year}.md"
            path = os.path.join(out_dir, file_name)
            if os.path.exists(path):
                skipped += 1
                continue
            note = build_note(name, args.account, args.category, year, projected, month_values)
            with open(path, "w", encoding="utf-8") as f:
                f.write(note)
            written += 1

    print(f"Wrote {written} notes to {out_dir} ({skipped} skipped).")


if __name__ == "__main__":
    main()
